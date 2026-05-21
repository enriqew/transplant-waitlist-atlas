"""Silver prebuild: ANZDATA Chapter 6 Excel files -> long-format CSV.

Each `anzdata_ch6_{data_year}.xlsx` file contains Table 6.1 (sheet `tab6_1_all`)
with a 6-year rolling window of Australian kidney transplant waiting list dynamics.
The "Active end of year" row gives the year-end waitlist stock.

Multiple Excel files overlap in calendar years. When the same calendar year
appears in more than one file, the value from the most recent report
(highest data_year) is used — this captures any late corrections.

Output schema:
    metric_type,country_iso3,year,organ,patients
    stock,AUS,2016,kidney,953
    stock,AUS,2017,kidney,965
    ...

Coverage: Australia (AUS) only, kidney only.
Year: calendar year-end (31 December), consistent with ET and CENATRA.
"""

from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[2]
RAW_ROOT = REPO_ROOT / "data" / "raw" / "anzdata_waitlist"

# Regex to extract data_year from filename: anzdata_ch6_{data_year}.xlsx
_DATA_YEAR_RE = re.compile(r"anzdata_ch6_(\d{4})\.xlsx", re.IGNORECASE)

# Sheet name for Table 6.1 (waitlist dynamics).
_SHEET_NAME = "tab6_1_all"

# Row label for year-end active waitlist count.
_ACTIVE_END_LABEL = "Active end of year"


def _latest_snapshot_dir() -> Path:
    if not RAW_ROOT.exists():
        raise SystemExit(f"no ANZDATA bronze snapshots under {RAW_ROOT}")
    candidates = sorted(p for p in RAW_ROOT.iterdir() if p.is_dir())
    if not candidates:
        raise SystemExit(f"no dated subdirectories under {RAW_ROOT}")
    for candidate in reversed(candidates):
        if any(candidate.glob("anzdata_ch6_*.xlsx")):
            return candidate
    raise SystemExit(
        f"no ANZDATA Chapter 6 xlsx files found under any snapshot in {RAW_ROOT}; "
        "run `python -m ingest.anzdata_waitlist` to download"
    )


def _parse_chapter6(xlsx_path: Path) -> dict[int, int]:
    """Parse Table 6.1 and return {calendar_year: active_end_of_year}."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if _SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"sheet '{_SHEET_NAME}' not found in {xlsx_path.name}")

    ws = wb[_SHEET_NAME]
    rows = list(ws.iter_rows(values_only=True))

    year_cols: list[int] = []
    results: dict[int, int] = {}

    for row in rows:
        if not row or all(c is None for c in row):
            continue
        first = str(row[0]).strip() if row[0] is not None else ""

        # Header row: contains year values (int or numeric string like '2019').
        if not year_cols:
            candidate_years = []
            for col_idx, cell in enumerate(row):
                if cell is None:
                    continue
                try:
                    y = int(str(cell).strip().replace(",", ""))
                    if 2000 <= y <= 2100:
                        candidate_years.append((col_idx, y))
                except (ValueError, TypeError):
                    pass
            if len(candidate_years) >= 4:
                year_cols = candidate_years
                continue

        # Data row: "Active end of year".
        if first.lower().startswith("active end"):
            for col_idx, year in year_cols:
                val = row[col_idx] if col_idx < len(row) else None
                if val is None:
                    continue
                try:
                    patients = int(str(val).replace(",", "").strip())
                    if patients > 0:
                        results[year] = patients
                except (ValueError, TypeError):
                    pass
            break  # Found the row we need; stop scanning.

    return results


def main() -> int:
    snapshot = _latest_snapshot_dir()
    xlsx_files = sorted(snapshot.glob("anzdata_ch6_*.xlsx"))
    if not xlsx_files:
        print(f"ERROR: no anzdata_ch6_*.xlsx files in {snapshot}", file=sys.stderr)
        return 1

    # Merge all files; most-recent report (highest data_year) wins per year.
    # {calendar_year: (data_year, patients)}
    merged: dict[int, tuple[int, int]] = {}

    for xlsx_path in xlsx_files:
        m = _DATA_YEAR_RE.match(xlsx_path.name)
        if not m:
            print(f"WARNING: cannot infer data_year from {xlsx_path.name}", file=sys.stderr)
            continue
        data_year = int(m.group(1))

        try:
            year_map = _parse_chapter6(xlsx_path)
        except Exception as exc:  # noqa: BLE001
            print(f"WARNING: failed to parse {xlsx_path.name}: {exc}", file=sys.stderr)
            continue

        for calendar_year, patients in year_map.items():
            existing = merged.get(calendar_year)
            if existing is None or data_year > existing[0]:
                merged[calendar_year] = (data_year, patients)

        print(f"  {xlsx_path.name} (data_year={data_year}): years {sorted(year_map)}")

    if not merged:
        print("ERROR: no data extracted from any ANZDATA xlsx", file=sys.stderr)
        return 1

    out_path = snapshot / "anzdata_long.csv"
    total_rows = 0

    with out_path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(["metric_type", "country_iso3", "year", "organ", "patients"])
        for calendar_year, (_, patients) in sorted(merged.items()):
            writer.writerow(["stock", "AUS", calendar_year, "kidney", patients])
            total_rows += 1

    print(f"wrote {total_rows:,} rows -> {out_path}")
    return 0 if total_rows > 0 else 1


if __name__ == "__main__":
    sys.exit(main())
